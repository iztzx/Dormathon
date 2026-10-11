"""Independent, offline backend contracts; all fixtures and training labels are synthetic."""

import asyncio
import csv
import hashlib
import json
import pickle
from copy import deepcopy
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import numpy as np
import pandas as pd
import pytest
import sklearn
from cryptography.exceptions import InvalidSignature
from cryptography.hazmat.primitives import serialization
from sklearn.dummy import DummyClassifier, DummyRegressor

from agent_service import orchestrator
from agent_service.sdk_runtime import SDKRuntime
from agent_service.store import Store
from tools import artifact_tool, ml_underwrite_tool
from tools.decision_memo_tool import build_decision_memo, expected_figures
from tools.ml_underwrite_tool import FEATURES, evaluate_credit_ml
from tools.selfcheck_tool import verify_bundle_integrity
from tools.reconciliation_tool import (
    EvidenceError,
    ParityViolation,
    enforce_parity,
    run_excel_reconciliation,
)


BANK_COLUMNS = ("transaction_id", "invoice_id", "currency", "inflow", "outflow", "date")
EXPECTED_FEATURES = [
    "monthly_inflow", "monthly_outflow", "monthly_debt_service", "invoice_total",
    "cashflow_margin", "dsr",
]


def synthetic_invoice(invoice_id="SYNTHETIC-INV-1", quantity="10", unit_price="10.00",
                      subtotal="100.00", total="100.00"):
    return {
        "invoice_id": invoice_id, "currency": "MYR", "subtotal": subtotal, "total": total,
        "tax": "0.00", "freight": "0.00", "discount": "0.00",
        "container_seal": "SYNTHETIC-SEAL-1", "incoterms": "FOB",
        "source": "synthetic-invoice.pdf", "regions": [{"page": 1, "bbox": [0, 0, 100, 100]}],
        "line_items": [{"sku": "SYNTHETIC-SKU-1", "quantity": quantity, "unit_price": unit_price}],
    }


def synthetic_delivery(invoice):
    return {
        "invoice_id": invoice["invoice_id"], "container_seal": invoice["container_seal"],
        "incoterms": invoice["incoterms"], "source": "synthetic-delivery.pdf",
        "line_items": [{"sku": item["sku"], "quantity": item["quantity"]}
                       for item in invoice["line_items"]],
    }


def synthetic_bank_row(invoice_id="SYNTHETIC-INV-1", inflow="100.00", transaction_id="SYNTHETIC-TX-1",
                       outflow="0.00", **changes):
    return {
        "transaction_id": transaction_id, "invoice_id": invoice_id, "currency": "MYR",
        "inflow": inflow, "outflow": outflow, "date": "2026-01-15", **changes,
    }


def write_statement(path, rows):
    if path.suffix == ".xlsx":
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = "BankLedger"
        sheet.append(BANK_COLUMNS)
        for row in rows:
            sheet.append([row[column] for column in BANK_COLUMNS])
        book.save(path)
        book.close()
    else:
        with path.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=BANK_COLUMNS)
            writer.writeheader()
            writer.writerows(rows)
    return path


@pytest.fixture
def valid_report(tmp_path):
    invoice = synthetic_invoice()
    statement = write_statement(tmp_path / "synthetic-bank.csv", [synthetic_bank_row()])
    return run_excel_reconciliation([invoice], str(statement), [synthetic_delivery(invoice)])


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_exact_quantity_times_unit_price_and_bank_parity(tmp_path, extension):
    invoice = synthetic_invoice(quantity="3", unit_price="0.10", subtotal="0.30", total="0.30")
    statement = write_statement(tmp_path / f"synthetic-bank.{extension}",
                                [synthetic_bank_row(inflow="0.30")])
    report = run_excel_reconciliation([invoice], str(statement), [synthetic_delivery(invoice)])
    assert report["passed"] is True
    assert Decimal(report["variance"]) == Decimal(report["ledger_delta"]) == 0
    assert Decimal(report["invoice_total"]) == Decimal(report["verified_inflow"]) == Decimal("0.30")
    assert report["errors"] == []
    assert report["gates"] == {"exact_money_parity": True, "delivery_verified": True}
    assert report["statement_sha256"] == hashlib.sha256(statement.read_bytes()).hexdigest()
    assert enforce_parity(report) is report


@pytest.mark.parametrize("gate", ["line_items_to_subtotal", "subtotal_adjustments_to_total", "invoice_to_bank"])
def test_equal_grand_totals_cannot_cancel_per_invoice_mismatches(tmp_path, gate):
    first = synthetic_invoice()
    second = synthetic_invoice(invoice_id="SYNTHETIC-INV-2")
    rows = [synthetic_bank_row(inflow="100.00"),
            synthetic_bank_row(second["invoice_id"], "100.00", "SYNTHETIC-TX-2")]
    if gate == "line_items_to_subtotal":
        first["line_items"][0]["unit_price"] = "9.00"
        second["line_items"][0]["unit_price"] = "11.00"
    elif gate == "subtotal_adjustments_to_total":
        first["subtotal"] = "90.00"
        first["line_items"][0]["unit_price"] = "9.00"
        second["subtotal"] = "110.00"
        second["line_items"][0]["unit_price"] = "11.00"
    else:
        rows[0]["inflow"], rows[1]["inflow"] = "90.00", "110.00"
    statement = write_statement(tmp_path / "synthetic-bank.csv", rows)
    report = run_excel_reconciliation([first, second], str(statement))
    errors = [error for error in report["errors"] if error["gate"] == gate]
    assert len(errors) == 2
    assert {Decimal(error["delta"]) for error in errors} == {Decimal("-10"), Decimal("10")}
    assert Decimal(report["ledger_delta"]) == 0
    assert Decimal(report["variance"]) == 10
    assert report["passed"] is False
    with pytest.raises(ParityViolation) as rejected:
        enforce_parity(report)
    assert rejected.value.report is report


@pytest.mark.parametrize("delta", ["0.01", "0.000001", "0.000000000001"])
def test_one_cent_and_fractional_line_deltas_never_round_to_zero(tmp_path, delta):
    invoice = synthetic_invoice(quantity="1", unit_price=str(Decimal("100") - Decimal(delta)))
    statement = write_statement(tmp_path / "synthetic-bank.csv", [synthetic_bank_row()])
    report = run_excel_reconciliation([invoice], str(statement))
    assert Decimal(report["variance"]) == Decimal(delta) != 0
    assert Decimal(report["ledger_delta"]) == 0
    assert report["passed"] is False
    assert report["errors"][0]["gate"] == "line_items_to_subtotal"
    assert Decimal(report["errors"][0]["delta"]) == -Decimal(delta)
    with pytest.raises(ParityViolation):
        enforce_parity(report)


def test_one_cent_bank_delta_is_nonzero(tmp_path):
    statement = write_statement(tmp_path / "synthetic-bank.csv", [synthetic_bank_row(inflow="99.99")])
    report = run_excel_reconciliation([synthetic_invoice()], str(statement))
    assert Decimal(report["variance"]) == Decimal(report["ledger_delta"]) == Decimal("0.01")
    with pytest.raises(ParityViolation):
        enforce_parity(report)


@pytest.mark.parametrize("unmatched_id", ["", "SYNTHETIC-UNKNOWN"])
def test_unmatched_inflows_are_not_available_to_cancel_missing_invoice_payment(tmp_path, unmatched_id):
    statement = write_statement(tmp_path / "synthetic-bank.csv", [synthetic_bank_row(unmatched_id)])
    report = run_excel_reconciliation([synthetic_invoice()], str(statement))
    assert Decimal(report["ledger_delta"]) == 0
    assert {error["gate"] for error in report["errors"]} == {"unmatched_bank_inflow", "invoice_to_bank"}
    assert Decimal(report["variance"]) == 100
    with pytest.raises(ParityViolation):
        enforce_parity(report)


def test_duplicate_invoice_ids_rejected(tmp_path):
    statement = write_statement(tmp_path / "synthetic-bank.csv", [synthetic_bank_row()])
    invoice = synthetic_invoice()
    with pytest.raises(EvidenceError, match="duplicate invoice_id"):
        run_excel_reconciliation([invoice, deepcopy(invoice)], str(statement))


def test_duplicate_transaction_ids_rejected_even_if_split_payments_sum_exactly(tmp_path):
    statement = write_statement(tmp_path / "synthetic-bank.csv",
                                [synthetic_bank_row(inflow="50.00"), synthetic_bank_row(inflow="50.00")])
    with pytest.raises(EvidenceError, match="duplicate transaction_id"):
        run_excel_reconciliation([synthetic_invoice()], str(statement))


def test_unique_split_payments_and_unlinked_outflows_are_permitted(tmp_path):
    rows = [synthetic_bank_row(inflow="40.00"), synthetic_bank_row(inflow="60.00", transaction_id="TX-2"),
            synthetic_bank_row("", "0.00", "TX-3", outflow="30.00")]
    statement = write_statement(tmp_path / "synthetic-bank.csv", rows)
    report = run_excel_reconciliation([synthetic_invoice()], str(statement))
    assert enforce_parity(report)["passed"] is True
    assert Decimal(report["verified_inflow"]) == 100
    assert len(report["bank_rows"]) == 3


@pytest.mark.parametrize("where", ["invoice", "bank"])
@pytest.mark.parametrize("currency", ["USD", "", "myr"])
def test_unverified_currency_rejected(tmp_path, where, currency):
    invoice, row = synthetic_invoice(), synthetic_bank_row()
    (invoice if where == "invoice" else row)["currency"] = currency
    statement = write_statement(tmp_path / "synthetic-bank.csv", [row])
    with pytest.raises(EvidenceError, match="currency|MYR"):
        run_excel_reconciliation([invoice], str(statement))


@pytest.mark.parametrize("quantity", ["-1", "0"])
def test_nonpositive_invoice_quantity_rejected(tmp_path, quantity):
    statement = write_statement(tmp_path / "synthetic-bank.csv", [synthetic_bank_row()])
    with pytest.raises(EvidenceError, match="positive quantity"):
        run_excel_reconciliation([synthetic_invoice(quantity=quantity)], str(statement))


@pytest.mark.parametrize("inflow,outflow", [("-1.00", "0.00"), ("100.00", "-1.00"), ("100.00", "1.00")])
def test_invalid_bank_amount_direction_rejected(tmp_path, inflow, outflow):
    statement = write_statement(tmp_path / "synthetic-bank.csv",
                                [synthetic_bank_row(inflow=inflow, outflow=outflow)])
    with pytest.raises(EvidenceError, match="direction"):
        run_excel_reconciliation([synthetic_invoice()], str(statement))


@pytest.mark.parametrize("field", ["inflow", "outflow", "invoice_id", "date"])
def test_formula_bank_xlsx_rejected_in_any_evidence_column(tmp_path, field):
    row = synthetic_bank_row()
    row[field] = "=1+1"
    statement = write_statement(tmp_path / "synthetic-bank.xlsx", [row])
    with pytest.raises(EvidenceError, match="formulas"):
        run_excel_reconciliation([synthetic_invoice()], str(statement))


@pytest.mark.parametrize("where", ["bank", "subtotal", "total"])
def test_subcent_money_rejected_without_rounding(tmp_path, where):
    invoice, row = synthetic_invoice(), synthetic_bank_row()
    (row if where == "bank" else invoice)["inflow" if where == "bank" else where] = "100.001"
    statement = write_statement(tmp_path / "synthetic-bank.csv", [row])
    with pytest.raises(EvidenceError, match="sub-cent"):
        run_excel_reconciliation([invoice], str(statement))


@pytest.mark.parametrize("defect,gate", [
    ("short_quantity", "delivery_quantity"), ("extra_sku", "delivery_quantity"),
    ("seal", "shipping_identity"), ("incoterms", "shipping_identity"),
    ("missing", "missing_delivery_docket"),
])
def test_packing_quantity_and_shipping_identity_are_independent_payout_gates(tmp_path, defect, gate):
    invoice = synthetic_invoice()
    delivery = synthetic_delivery(invoice)
    if defect == "short_quantity":
        delivery["line_items"][0]["quantity"] = "9.999999999999"
    elif defect == "extra_sku":
        delivery["line_items"].append({"sku": "SYNTHETIC-EXTRA", "quantity": "1"})
    elif defect == "seal":
        delivery["container_seal"] = "SYNTHETIC-WRONG-SEAL"
    elif defect == "incoterms":
        delivery["incoterms"] = "CIF"
    statement = write_statement(tmp_path / "synthetic-bank.csv", [synthetic_bank_row()])
    report = run_excel_reconciliation([invoice], str(statement), [] if defect == "missing" else [delivery])
    assert Decimal(report["variance"]) == Decimal(report["ledger_delta"]) == 0
    assert report["passed"] is False
    assert report["gates"]["delivery_verified"] is False
    assert gate in {error["gate"] for error in report["errors"]}
    with pytest.raises(ParityViolation):
        enforce_parity(report)


@pytest.mark.parametrize("defect", ["unknown_invoice", "negative_quantity"])
def test_invalid_delivery_evidence_rejected(tmp_path, defect):
    invoice = synthetic_invoice()
    delivery = synthetic_delivery(invoice)
    if defect == "unknown_invoice":
        delivery["invoice_id"] = "SYNTHETIC-UNKNOWN"
    else:
        delivery["line_items"][0]["quantity"] = "-1"
    statement = write_statement(tmp_path / "synthetic-bank.csv", [synthetic_bank_row()])
    with pytest.raises(EvidenceError):
        run_excel_reconciliation([invoice], str(statement), [delivery])


@pytest.mark.parametrize("field,value", [("passed", False), ("passed", 1), ("passed", "true"),
                                          ("variance", "0.01"), ("variance", "0.000000000001"),
                                          ("ledger_delta", "-0.01")])
def test_parity_guard_does_not_trust_truthiness_or_tolerances(valid_report, field, value):
    valid_report[field] = value
    with pytest.raises(ParityViolation):
        enforce_parity(valid_report)


@pytest.mark.parametrize("report", [{}, {"status": "ok"}, {"passed": True},
                                      {"passed": True, "variance": "NaN", "ledger_delta": "0"}])
def test_unknown_or_malformed_reports_fail_closed(report):
    with pytest.raises(EvidenceError):
        enforce_parity(report)


def test_parity_guard_rejects_incomplete_success_report():
    with pytest.raises(EvidenceError):
        enforce_parity({"passed": True, "variance": "0", "ledger_delta": "0"})


@pytest.mark.parametrize("contradiction", ["errors", "money_gate", "delivery_gate"])
def test_parity_guard_rejects_success_flag_contradicting_evidence(valid_report, contradiction):
    if contradiction == "errors":
        valid_report["errors"] = [{"gate": "invoice_to_bank", "delta": "0.01"}]
    else:
        valid_report["gates"]["exact_money_parity" if contradiction == "money_gate" else "delivery_verified"] = False
    with pytest.raises(EvidenceError):
        enforce_parity(valid_report)


@pytest.fixture
def runtime(tmp_path):
    job_dir = tmp_path / "job"
    normalized = job_dir / "normalized"
    normalized.mkdir(parents=True)
    (normalized / "synthetic-evidence.json").write_text("{}", encoding="utf-8")
    emit = Mock()

    async def dispatch(name, args):
        raise AssertionError("A hook test must not dispatch an actual tool")

    return SDKRuntime(job_dir, emit, dispatch)


@pytest.mark.parametrize("absolute", [False, True])
def test_read_hook_allows_only_current_normalized_json(runtime, absolute):
    path = "normalized/synthetic-evidence.json"
    if absolute:
        path = str(runtime.job_dir / path)
    result = asyncio.run(runtime.pre_tool({"tool_name": "Read", "tool_input": {"file_path": path}}))
    assert result == {}
    assert runtime.tool_calls == 1
    runtime.emit.assert_not_called()


@pytest.mark.parametrize("path", [
    "normalized/../../private.json", "normalized/../inputs/source.json",
    "../other-job/normalized/evidence.json", "normalized-evil/evidence.json",
    "normalized/private.pem", "inputs/synthetic-invoice.pdf", "normalized/evidence.JSON",
])
def test_read_hook_blocks_source_traversal_and_non_evidence_paths(runtime, path):
    result = asyncio.run(runtime.pre_tool({"tool_name": "Read", "tool_input": {"file_path": path}}))
    assert result["decision"] == "block"
    assert runtime.tool_calls == 0
    assert runtime.emit.call_args.args[1]["status"] == "BLOCKED"


def test_read_hook_blocks_absolute_other_job_path(runtime, tmp_path):
    result = asyncio.run(runtime.pre_tool({"tool_name": "Read", "tool_input": {
        "file_path": str(tmp_path / "other-job" / "normalized" / "evidence.json")}}))
    assert result["decision"] == "block"


@pytest.mark.parametrize("name", [
    "Bash", "PowerShell", "Shell", "curl", "WebFetch", "WebSearch", "Write", "Edit",
    "mcp__external__curl", "mcp__underwriting__run_excel_reconciliation",
    "mcp__underwriting__evaluate_credit_ml", "mcp__underwriting__compile_tnc_agreement", "UnknownTool",
])
def test_pre_hook_blocks_shell_curl_network_and_later_stage_tools(runtime, name):
    result = asyncio.run(runtime.pre_tool({"tool_name": name, "tool_input": {"command": "curl --version"}}))
    assert result["decision"] == "block"
    assert runtime.tool_calls == 0


@pytest.mark.parametrize("name", ["mcp__underwriting__run_ocr_normalizer", "mcp__underwriting__submit_trade_proposal"])
def test_pre_hook_permits_extraction_only_tools(runtime, name):
    assert asyncio.run(runtime.pre_tool({"tool_name": name, "tool_input": {}})) == {}


def test_tool_budget_blocks_twenty_fifth_authorized_call(runtime):
    payload = {"tool_name": "Read", "tool_input": {"file_path": "normalized/synthetic-evidence.json"}}
    for _ in range(24):
        assert asyncio.run(runtime.pre_tool(payload)) == {}
    assert asyncio.run(runtime.pre_tool(payload))["decision"] == "block"
    assert runtime.tool_calls == 25


@pytest.mark.parametrize("wrapped", [False, True])
def test_post_hook_parity_failure_triggers_sticky_rollback(runtime, valid_report, wrapped):
    bad = deepcopy(valid_report)
    bad["variance"] = "0.000000000001"
    output = {"content": [{"type": "text", "text": json.dumps(bad)}]} if wrapped else bad
    result = asyncio.run(runtime.post_tool({"tool_name": "mcp__underwriting__run_excel_reconciliation",
                                            "tool_response": output}))
    assert result["decision"] == "block"
    assert result["continue_"] is False
    assert runtime.rollback is True
    assert asyncio.run(runtime.post_tool({"tool_name": "run_excel_reconciliation",
                                          "tool_response": valid_report})) == {}
    assert runtime.rollback is True


@pytest.mark.parametrize("wrapped", [False, True])
def test_post_hook_accepts_exact_valid_report(runtime, valid_report, wrapped):
    output = {"content": [{"type": "text", "text": json.dumps(valid_report)}]} if wrapped else valid_report
    assert asyncio.run(runtime.post_tool({"tool_name": "run_excel_reconciliation", "tool_response": output})) == {}
    assert runtime.rollback is False


@pytest.mark.parametrize("output", [
    {}, {"status": "unknown"}, {"content": []}, {"content": [{"type": "image"}]},
    {"content": [{"type": "text", "text": "not JSON"}]},
    {"content": [{"type": "text", "text": "null"}]},
    {"content": [{"type": "text", "text": "[]"}]}, None, [], "unknown report",
], ids=["empty-report", "unknown-schema", "empty-content", "no-text-content", "invalid-json",
        "json-null", "json-array", "null-response", "array-response", "string-response"])
def test_post_hook_unknown_or_malformed_response_blocks_and_rolls_back(runtime, output):
    result = asyncio.run(runtime.post_tool({"tool_name": "run_excel_reconciliation", "tool_response": output}))
    assert result["decision"] == "block"
    assert result["continue_"] is False
    assert runtime.rollback is True


@pytest.fixture
def store(tmp_path):
    return Store(tmp_path / "workspace")


def wallet_row(store, jid):
    with store.connect() as db:
        row = db.execute("SELECT amount,status FROM wallet WHERE job_id=?", (jid,)).fetchone()
    return dict(row) if row is not None else None


@pytest.mark.parametrize("destination", ["UNDERWRITING", "SETTLING", "COMPLETED", "REVIEW_REQUIRED", "UNKNOWN"])
def test_forbidden_state_transition_does_not_change_state_or_events(store, destination):
    jid, _ = store.create("demo", {}, [])
    before = store.events(jid)
    with pytest.raises(ValueError, match="Invalid transition"):
        store.transition(jid, destination)
    assert store.get(jid)["state"] == "QUEUED"
    assert store.events(jid) == before
    assert wallet_row(store, jid) is None


def test_review_required_cannot_transition_to_settling(store):
    jid, _ = store.create("demo", {}, [])
    for state in ("EXTRACTING", "RECONCILING", "UNDERWRITING", "REVIEW_REQUIRED"):
        store.transition(jid, state)
    with pytest.raises(ValueError, match="Invalid transition"):
        store.transition(jid, "SETTLING")
    assert store.get(jid)["state"] == "REVIEW_REQUIRED"
    assert wallet_row(store, jid) is None


def test_stream_tokens_are_job_scoped_and_only_hashes_persist(store):
    first, first_token = store.create("demo", {}, [])
    second, second_token = store.create("demo", {}, [])
    assert first_token != second_token
    assert store.token_ok(first, first_token)
    assert store.token_ok(second, second_token)
    assert not store.token_ok(first, second_token)
    assert not store.token_ok(second, first_token)
    assert not store.token_ok(first, "")
    assert not store.token_ok("0" * 32, first_token)
    assert not store.token_ok("' OR 1=1 --", first_token)
    with store.connect() as db:
        row = db.execute("SELECT token_hash,data FROM jobs WHERE id=?", (first,)).fetchone()
    assert row["token_hash"] == hashlib.sha256(first_token.encode()).hexdigest()
    assert first_token not in row["data"]
    assert "token_hash" not in store.get(first)


@pytest.mark.parametrize("state", ["QUEUED", "EXTRACTING", "RECONCILING", "SELF_HEALING", "UNDERWRITING", "SETTLING"])
def test_recover_interrupted_jobs_fails_closed_and_rolls_back_wallet(store, state):
    jid, token = store.create("demo", {}, [])
    store.update(jid, state=state, metrics={"synthetic": True})
    store.settle_demo(jid, "100.00")
    restarted = Store(store.workspace)
    restarted.recover()
    job = restarted.get(jid)
    assert job["state"] == "FAILED"
    assert "interrupted" in job["error"].lower()
    assert job["metrics"] == {"synthetic": True}
    assert wallet_row(restarted, jid) == {"amount": "100.00", "status": "ROLLED_BACK"}
    assert restarted.token_ok(jid, token)
    events = restarted.events(jid)
    restarted.recover()
    assert restarted.events(jid) == events


@pytest.mark.parametrize("state", ["COMPLETED", "REVIEW_REQUIRED", "QUARANTINED", "FAILED"])
def test_recovery_preserves_terminal_jobs(store, state):
    jid, _ = store.create("demo", {}, [])
    store.update(jid, state=state)
    before = store.get(jid), store.events(jid)
    Store(store.workspace).recover()
    assert (store.get(jid), store.events(jid)) == before
    assert wallet_row(store, jid) is None


def test_wallet_simulation_is_idempotent_and_rejects_amount_changes(store):
    jid, _ = store.create("demo", {}, [])
    store.settle_demo(jid, "100.00")
    store.settle_demo(jid, "100.00")
    with pytest.raises(ValueError, match="Idempotency conflict"):
        store.settle_demo(jid, "100.01")
    assert wallet_row(store, jid) == {"amount": "100.00", "status": "PREPARED_SIMULATION"}
    store.rollback_wallet(jid)
    assert wallet_row(store, jid)["status"] == "ROLLED_BACK"


def test_event_hash_chain_is_independently_verifiable_and_job_scoped(store):
    first, _ = store.create("demo", {}, [])
    second, _ = store.create("demo", {}, [])
    store.event(first, "SYNTHETIC_AUDIT", {"variance": "0", "passed": True})
    store.event(second, "SYNTHETIC_AUDIT", {"variance": "0.01"})
    previous = "0" * 64
    for event in store.events(first):
        encoded = json.dumps(event["payload"], sort_keys=True, separators=(",", ":"), allow_nan=False)
        digest = hashlib.sha256((previous + first + event["type"] + event["timestamp"] + encoded).encode()).hexdigest()
        assert event["job_id"] == first
        assert event["previous_hash"] == previous
        assert event["sha256"] == digest
        previous = digest
    assert store.events(second)[0]["previous_hash"] == "0" * 64


@pytest.fixture
def engine_factory(tmp_path, monkeypatch):
    def create(*, mismatch=False, mode="demo"):
        store = Store(tmp_path / "engine-workspace")
        invoice = synthetic_invoice(quantity="1000", unit_price="100.00", subtotal="100000.00", total="100000.00")
        delivery = synthetic_delivery(invoice)
        if mismatch:
            invoice["line_items"][0]["unit_price"] = "99.99999"
        terms = {"requested_limit": "40000.00", "duration_months": 12,
                 "annual_rate": "0.08", "monthly_debt_service": "1000.00"}
        jid, token = store.create(mode, terms, [])
        directory = store.workspace / "jobs" / jid
        inputs = directory / "inputs"
        inputs.mkdir(parents=True)
        # These are byte fixtures only; OCR is mocked, not claimed to have read real PDFs.
        for name in (invoice["source"], delivery["source"]):
            (inputs / name).write_bytes(b"synthetic unit-test evidence; not an actual PDF")
        statement = write_statement(inputs / "synthetic-bank.csv", [
            synthetic_bank_row(inflow="100000.00"),
            synthetic_bank_row("", "0.00", "SYNTHETIC-TX-2", outflow="30000.00"),
        ])
        sources = [{"name": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
                   for path in sorted(inputs.iterdir())]
        store.update(jid, sources=sources)

        def normalize(path, regions=None, cycle=0):
            return {"source": Path(path).name, "cycle": cycle, "crops": []}

        def parse(packet):
            is_invoice = packet["source"] == invoice["source"]
            return {"kind": "invoice" if is_invoice else "delivery",
                    "data": deepcopy(invoice if is_invoice else delivery)}

        ocr = Mock(side_effect=normalize)
        risk = Mock(return_value={"passed": True, "review_required": False, "approved_limit": "40000.00",
                                  "proposed_limit": "40000.00", "default_probability": "0.0"})
        artifacts = Mock(return_value=["synthetic-bundle.bin"])
        selfcheck = Mock(return_value={"verified": True, "outputs_checked": 7, "sources_checked": 10,
                                       "self_check": "post_commit", "is_demo": True,
                                       "notice": "Synthetic mock; real bundles verify in integration tests"})
        monkeypatch.setattr(orchestrator, "run_ocr_normalizer", ocr)
        monkeypatch.setattr(orchestrator, "parse_demo_packet", parse)
        monkeypatch.setattr(orchestrator, "evaluate_credit_ml", risk)
        monkeypatch.setattr(orchestrator, "build_artifacts", artifacts)
        monkeypatch.setattr(orchestrator, "verify_bundle_integrity", selfcheck)
        return SimpleNamespace(store=store, jid=jid, token=token, directory=directory, statement=statement,
                               invoice=invoice, delivery=delivery, ocr=ocr, risk=risk, artifacts=artifacts,
                               selfcheck=selfcheck, engine=orchestrator.UnderwritingEngine(store))
    return create


def test_engine_quarantines_after_exactly_three_self_heals_without_underwriting_or_payout(engine_factory):
    case = engine_factory(mismatch=True)
    asyncio.run(case.engine.run(case.jid))
    job = case.store.get(case.jid)
    assert job["state"] == "QUARANTINED"
    assert Decimal(job["metrics"]["variance"]) == Decimal("0.01")
    assert [call.args[2] for call in case.ocr.call_args_list] == [0, 0, 1, 1, 2, 2, 3, 3]
    events = case.store.events(case.jid)
    states = [event["payload"]["state"] for event in events if event["type"] == "STATE_TRANSITION"]
    assert states.count("SELF_HEALING") == 3
    assert states.count("RECONCILING") == 4
    assert "UNDERWRITING" not in states and "SETTLING" not in states
    rollback = [event for event in events if event["type"] == "ROLLBACK_TRIGGERED"]
    assert len(rollback) == 4
    assert all(event["payload"]["maximum_self_heals"] == 3 for event in rollback)
    report = json.loads((case.directory / "quarantine" / "variance_report.json").read_text(encoding="utf-8"))
    assert report["self_heal_cycles"] == 3
    assert report["disbursement"] == "prohibited"
    assert report["passed"] is False
    case.risk.assert_not_called()
    case.artifacts.assert_not_called()
    case.selfcheck.assert_not_called()
    assert wallet_row(case.store, case.jid) is None
    assert not (case.directory / "artifacts").exists()


@pytest.mark.parametrize("heal_cycle", [1, 3])
def test_engine_successful_self_heal_stops_at_first_exact_audit(engine_factory, monkeypatch, heal_cycle):
    case = engine_factory()
    real = run_excel_reconciliation
    count = 0

    def reconcile(invoices, statement_path, deliveries):
        nonlocal count
        records = deepcopy(invoices)
        if count < heal_cycle:
            records[0]["line_items"][0]["unit_price"] = "99.99999"
        count += 1
        return real(records, statement_path, deliveries)

    monkeypatch.setattr(orchestrator, "run_excel_reconciliation", reconcile)
    asyncio.run(case.engine.run(case.jid))
    assert count == heal_cycle + 1
    assert case.ocr.call_count == 2 * (heal_cycle + 1)
    assert case.store.get(case.jid)["state"] == "COMPLETED"
    assert Decimal(case.store.get(case.jid)["metrics"]["variance"]) == 0
    case.risk.assert_called_once()
    case.artifacts.assert_called_once()
    case.selfcheck.assert_called_once()
    assert wallet_row(case.store, case.jid) == {"amount": "40000.00", "status": "SIMULATED_ONLY"}


def test_engine_risk_review_never_compiles_artifacts_or_prepares_wallet(engine_factory):
    case = engine_factory()
    case.risk.return_value = {"passed": False, "review_required": True, "approved_limit": "0.00",
                              "proposed_limit": "20000.00", "haircut_requested_limit": "20000.00"}
    asyncio.run(case.engine.run(case.jid))
    job = case.store.get(case.jid)
    assert job["state"] == "REVIEW_REQUIRED"
    assert job["metrics"]["approved_limit"] == "0.00"
    assert job["artifacts"] == []
    assert job["decision_memo"]["outcome"] == "REVIEW_REQUIRED"
    assert job["decision_memo"]["is_demo"] is True
    case.artifacts.assert_not_called()
    case.selfcheck.assert_not_called()
    assert wallet_row(case.store, case.jid) is None
    assert not (case.directory / "artifacts").exists()


def test_engine_checks_final_parity_again_before_settlement(engine_factory, monkeypatch):
    case = engine_factory()
    report = run_excel_reconciliation([case.invoice], str(case.statement), [case.delivery])
    monkeypatch.setattr(orchestrator, "run_excel_reconciliation", Mock(return_value=report))

    def invalidate_during_risk(*args, **kwargs):
        report["variance"] = "0.01"
        report["passed"] = False
        return {"passed": True, "approved_limit": "40000.00"}

    case.risk.side_effect = invalidate_during_risk
    asyncio.run(case.engine.run(case.jid))
    assert case.store.get(case.jid)["state"] == "FAILED"
    case.artifacts.assert_not_called()
    assert wallet_row(case.store, case.jid) is None


def test_engine_artifact_failure_rolls_back_prepared_wallet_and_hides_artifacts(engine_factory):
    case = engine_factory()
    case.artifacts.side_effect = EvidenceError("Synthetic signature-stage failure")
    asyncio.run(case.engine.run(case.jid))
    job = case.store.get(case.jid)
    assert job["state"] == "FAILED"
    assert job["artifacts"] == []
    assert job["error"] == "Synthetic signature-stage failure"
    assert wallet_row(case.store, case.jid) == {"amount": "40000.00", "status": "ROLLED_BACK"}


def test_engine_source_hash_tampering_fails_before_ocr(engine_factory):
    case = engine_factory()
    (case.directory / "inputs" / case.invoice["source"]).write_bytes(b"changed synthetic evidence")
    asyncio.run(case.engine.run(case.jid))
    assert case.store.get(case.jid)["state"] == "FAILED"
    assert "source hash" in case.store.get(case.jid)["error"].lower()
    case.ocr.assert_not_called()
    case.risk.assert_not_called()
    case.artifacts.assert_not_called()
    assert wallet_row(case.store, case.jid) is None


def test_engine_source_inventory_traversal_rejected_even_with_matching_hash(engine_factory):
    case = engine_factory()
    outside = case.directory / "outside.csv"
    outside.write_bytes(case.statement.read_bytes())
    sources = [source for source in case.store.get(case.jid)["sources"] if source["name"] != case.statement.name]
    sources.append({"name": "../outside.csv", "sha256": hashlib.sha256(outside.read_bytes()).hexdigest()})
    case.store.update(case.jid, sources=sources)
    asyncio.run(case.engine.run(case.jid))
    assert case.store.get(case.jid)["state"] == "FAILED"
    case.ocr.assert_not_called()
    case.risk.assert_not_called()
    case.artifacts.assert_not_called()
    assert wallet_row(case.store, case.jid) is None


@pytest.mark.parametrize("name,args", [
    ("run_ocr_normalizer", {"file_path": "../other-job/source.pdf"}),
    ("evaluate_credit_ml", {"financial_metrics": {}}),
    ("run_excel_reconciliation", {}),
])
def test_live_extraction_dispatch_blocks_path_traversal_and_stage_escalation(engine_factory, monkeypatch, name, args):
    case = engine_factory(mode="live")

    async def attempt_unauthorized_dispatch(self, paths, cycle=0, prior_errors=None):
        return await self.dispatch(name, args)

    monkeypatch.setattr(SDKRuntime, "extract", attempt_unauthorized_dispatch)
    asyncio.run(case.engine.run(case.jid))
    assert case.store.get(case.jid)["state"] == "FAILED"
    assert case.ocr.call_count == 2  # Only the two legitimate initial PDFs were normalized.
    case.risk.assert_not_called()
    case.artifacts.assert_not_called()
    assert wallet_row(case.store, case.jid) is None


def synthetic_metrics(debt="1000.00", requested="40000.00"):
    income, expense, debt_value = Decimal("100000.00"), Decimal("30000.00"), Decimal(debt)
    return {
        "monthly_inflow": str(income), "monthly_outflow": str(expense),
        "monthly_debt_service": str(debt_value), "invoice_total": "100000.00",
        "cashflow_margin": format((income - expense) / income, ".12f"),
        "dsr": format(debt_value / income, ".12f"), "requested_limit": requested,
        "provenance": "SYNTHETIC UNIT TEST ONLY; not actual borrower metrics or training labels",
    }


@pytest.fixture
def model_factory(tmp_path):
    counter = 0

    def create(*, constant=0, ceiling=50000, default_class=1):
        nonlocal counter
        counter += 1
        directory = tmp_path / f"synthetic-models-{counter}"
        directory.mkdir()
        # Both labels are explicit synthetic test classes, never actual credit/default observations.
        frame = pd.DataFrame([[100000, 30000, 1000, 100000, .7, .01],
                              [100000, 30000, 50000, 100000, .7, .5]], columns=FEATURES)
        classifier = DummyClassifier(strategy="constant", constant=constant).fit(frame, [0, 1])
        regressor = DummyRegressor(strategy="constant", constant=ceiling).fit(frame, [50000, 50000])
        hashes = {}
        for filename, model in (("credit_classifier.pkl", classifier), ("credit_regressor.pkl", regressor)):
            data = pickle.dumps(model, protocol=pickle.HIGHEST_PROTOCOL)
            (directory / filename).write_bytes(data)
            hashes[filename] = hashlib.sha256(data).hexdigest()
        manifest = {"is_demo": True, "sklearn_version": sklearn.__version__, "feature_names": list(FEATURES),
                    "default_class": default_class, "sha256": hashes,
                    "training_data": "synthetic unit-test rows and labels only"}
        (directory / "model_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
        return directory, manifest
    return create


def test_ml_feature_contract_is_explicit():
    assert FEATURES == EXPECTED_FEATURES
    metrics = synthetic_metrics()
    assert metrics["dsr"] == "0.010000000000"
    assert metrics["cashflow_margin"] == "0.700000000000"
    assert synthetic_metrics("50000.00")["dsr"] == "0.500000000000"


def test_synthetic_zero_default_probability_approves_positive_bounded_limit(model_factory):
    directory, manifest = model_factory(constant=0)
    result = evaluate_credit_ml(synthetic_metrics(requested="60000.00"), str(directory), mode="demo")
    assert result["passed"] is True and result["review_required"] is False
    assert Decimal(result["default_probability"]) == 0
    assert Decimal(result["financing_ceiling"]) == Decimal(result["approved_limit"]) == 50000
    assert 0 < Decimal(result["approved_limit"]) <= Decimal(result["requested_limit"])
    assert Decimal(result["approved_limit"]) <= Decimal("100000") * Decimal("0.8")
    assert Decimal(result["dsr"]) <= Decimal("0.45")
    assert result["haircut_requested_limit"] is None
    assert result["model_hashes"] == manifest["sha256"]
    assert result["model_manifest_sha256"] == hashlib.sha256((directory / "model_manifest.json").read_bytes()).hexdigest()
    assert result["is_demo_model"] is True
    assert "SYNTHETIC" in result["feature_provenance"]


def test_synthetic_default_probability_one_requires_review_with_exact_requested_haircut(model_factory):
    directory, _ = model_factory(constant=1)
    result = evaluate_credit_ml(synthetic_metrics(), str(directory), mode="demo", requested_limit="40000.00")
    assert Decimal(result["default_probability"]) == 1
    assert result["passed"] is False and result["review_required"] is True
    assert result["approved_limit"] == "0.00"
    assert Decimal(result["haircut_requested_limit"]) == Decimal("40000") * Decimal("0.50")
    assert Decimal(result["proposed_limit"]) == 20000
    assert result["gates"]["probability_le_0_12"] is False


def test_haircut_uses_requested_amount_not_lower_model_ceiling(model_factory):
    directory, _ = model_factory(constant=1, ceiling=50000)
    result = evaluate_credit_ml(synthetic_metrics(requested="120000.00"), str(directory), mode="demo")
    assert result["passed"] is False
    assert result["approved_limit"] == "0.00"
    assert Decimal(result["haircut_requested_limit"]) == 60000
    assert Decimal(result["proposed_limit"]) == 50000


def test_existing_debt_above_dsr_gate_never_approves_even_with_zero_default_probability(model_factory):
    directory, _ = model_factory(constant=0)
    result = evaluate_credit_ml(synthetic_metrics("50000.00"), str(directory), mode="demo")
    assert Decimal(result["default_probability"]) == 0
    assert Decimal(result["existing_dsr"]) == Decimal("0.5")
    assert result["passed"] is False
    assert result["review_required"] is True
    assert result["approved_limit"] == "0.00"
    assert Decimal(result["proposed_limit"]) == 0
    assert Decimal(result["haircut_requested_limit"]) == 20000
    assert result["gates"]["all_in_dsr_le_0_45"] is False


@pytest.mark.parametrize("constant,default_class,expected", [(0, 1, 0), (1, 1, 1), (0, 0, 1), (1, 0, 0)])
def test_default_probability_uses_manifest_label_not_guessed_column_zero(model_factory, constant, default_class, expected):
    directory, _ = model_factory(constant=constant, default_class=default_class)
    result = evaluate_credit_ml(synthetic_metrics(), str(directory), mode="demo")
    assert Decimal(result["default_probability"]) == expected
    assert result["passed"] is (expected == 0)


@pytest.mark.parametrize("probability,approved", [(0.12, True), (0.120000000001, False)])
def test_probability_point_twelve_boundary_is_inclusive(model_factory, monkeypatch, probability, approved):
    directory, _ = model_factory()
    # Patch the sklearn class after serializing real fitted DummyClassifier models;
    # no local callable is pickled and no production model is fabricated.
    monkeypatch.setattr(DummyClassifier, "predict_proba",
                        lambda self, frame: np.array([[1 - probability, probability]]))
    result = evaluate_credit_ml(synthetic_metrics(), str(directory), mode="demo")
    assert result["passed"] is approved
    assert result["gates"]["probability_le_0_12"] is approved
    if not approved:
        assert result["approved_limit"] == "0.00"


@pytest.mark.parametrize("output,value", [("probability", float("nan")), ("probability", float("inf")),
                                          ("probability", -0.1), ("probability", 1.1),
                                          ("ceiling", float("nan")), ("ceiling", float("inf")),
                                          ("ceiling", -1.0)])
def test_nonfinite_or_invalid_ml_outputs_fail_closed(model_factory, monkeypatch, output, value):
    directory, _ = model_factory()
    if output == "probability":
        monkeypatch.setattr(DummyClassifier, "predict_proba", lambda self, frame: np.array([[0, value]]))
    else:
        monkeypatch.setattr(DummyRegressor, "predict", lambda self, frame: np.array([value]))
    with pytest.raises(EvidenceError, match="nonfinite or invalid"):
        evaluate_credit_ml(synthetic_metrics(), str(directory), mode="demo")


@pytest.mark.parametrize("field", EXPECTED_FEATURES)
@pytest.mark.parametrize("value", ["NaN", "Infinity", "-Infinity"])
def test_nonfinite_financial_features_rejected_before_unpickle(model_factory, monkeypatch, field, value):
    directory, _ = model_factory()
    loader = Mock(side_effect=AssertionError("Unpickling unvalidated evidence is prohibited"))
    monkeypatch.setattr(ml_underwrite_tool.pickle, "load", loader)
    metrics = synthetic_metrics()
    metrics[field] = value
    with pytest.raises(EvidenceError):
        evaluate_credit_ml(metrics, str(directory), mode="demo")
    loader.assert_not_called()


@pytest.mark.parametrize("field", EXPECTED_FEATURES)
def test_missing_financial_feature_rejected_before_unpickle(model_factory, monkeypatch, field):
    directory, _ = model_factory()
    loader = Mock(side_effect=AssertionError("Missing feature must fail before unpickle"))
    monkeypatch.setattr(ml_underwrite_tool.pickle, "load", loader)
    metrics = synthetic_metrics()
    del metrics[field]
    with pytest.raises(EvidenceError):
        evaluate_credit_ml(metrics, str(directory), mode="demo")
    loader.assert_not_called()


def test_dsr_feature_must_match_synthetic_evidence_before_unpickle(model_factory, monkeypatch):
    directory, _ = model_factory()
    loader = Mock(side_effect=AssertionError("Inconsistent DSR must fail before unpickle"))
    monkeypatch.setattr(ml_underwrite_tool.pickle, "load", loader)
    metrics = synthetic_metrics()
    metrics["dsr"] = "0.020000000000"
    with pytest.raises(EvidenceError, match="DSR feature differs"):
        evaluate_credit_ml(metrics, str(directory), mode="demo")
    loader.assert_not_called()


@pytest.mark.parametrize("defect", ["missing_manifest", "malformed_json", "missing_hashes", "tampered_classifier",
                                     "tampered_regressor", "missing_model", "sklearn_version", "feature_order",
                                     "feature_missing", "feature_unknown"])
def test_manifest_schema_version_and_hash_checks_precede_all_unpickling(model_factory, monkeypatch, defect):
    directory, manifest = model_factory()
    manifest_path = directory / "model_manifest.json"
    if defect == "missing_manifest":
        manifest_path.unlink()
    elif defect == "malformed_json":
        manifest_path.write_text("not JSON", encoding="utf-8")
    elif defect.startswith("tampered_"):
        path = directory / ("credit_classifier.pkl" if defect == "tampered_classifier" else "credit_regressor.pkl")
        path.write_bytes(path.read_bytes() + b"synthetic tamper")
    elif defect == "missing_model":
        (directory / "credit_regressor.pkl").unlink()
    else:
        if defect == "missing_hashes":
            manifest.pop("sha256")
        elif defect == "sklearn_version":
            manifest["sklearn_version"] = "0.0.0-synthetic-mismatch"
        elif defect == "feature_order":
            manifest["feature_names"] = list(reversed(FEATURES))
        elif defect == "feature_missing":
            manifest["feature_names"] = FEATURES[:-1]
        elif defect == "feature_unknown":
            manifest["feature_names"] = FEATURES + ["synthetic_unknown"]
        manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    loader = Mock(side_effect=AssertionError("Unreviewed model was unpickled"))
    monkeypatch.setattr(ml_underwrite_tool.pickle, "load", loader)
    with pytest.raises((EvidenceError, json.JSONDecodeError)):
        evaluate_credit_ml(synthetic_metrics(), str(directory), mode="demo")
    loader.assert_not_called()


def test_synthetic_demo_models_cannot_run_in_default_live_mode(model_factory, monkeypatch):
    directory, _ = model_factory()
    loader = Mock(side_effect=AssertionError("Synthetic models cannot be unpickled for live inference"))
    monkeypatch.setattr(ml_underwrite_tool.pickle, "load", loader)
    with pytest.raises(EvidenceError, match="Synthetic demo models"):
        evaluate_credit_ml(synthetic_metrics(), str(directory))
    loader.assert_not_called()


@pytest.mark.parametrize("defect", ["unknown", "missing"])
def test_absent_manifest_default_class_is_rejected_before_inference(model_factory, monkeypatch, defect):
    directory, manifest = model_factory(default_class=2)
    if defect == "missing":
        manifest.pop("default_class")
        (directory / "model_manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    probability = Mock(side_effect=AssertionError("Default class must be checked before inference"))
    ceiling = Mock(side_effect=AssertionError("Default class must be checked before inference"))
    monkeypatch.setattr(DummyClassifier, "predict_proba", probability)
    monkeypatch.setattr(DummyRegressor, "predict", ceiling)
    with pytest.raises(EvidenceError, match="Default class label is absent"):
        evaluate_credit_ml(synthetic_metrics(), str(directory), mode="demo")
    probability.assert_not_called()
    ceiling.assert_not_called()


@pytest.mark.parametrize("model_type", ["classifier", "regressor"])
def test_loaded_model_feature_order_must_match_manifest(model_factory, monkeypatch, model_type):
    directory, _ = model_factory()
    filename = "credit_classifier.pkl" if model_type == "classifier" else "credit_regressor.pkl"
    # The contract is checked before prediction, so alter the deserialized model itself.
    with (directory / filename).open("rb") as stream:
        model = pickle.load(stream)
    model.feature_names_in_ = np.array(list(reversed(FEATURES)), dtype=object)
    data = pickle.dumps(model)
    (directory / filename).write_bytes(data)
    manifest_path = directory / "model_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["sha256"][filename] = hashlib.sha256(data).hexdigest()
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    probability = Mock(side_effect=AssertionError("Feature order must be checked before inference"))
    ceiling = Mock(side_effect=AssertionError("Feature order must be checked before inference"))
    monkeypatch.setattr(DummyClassifier, "predict_proba", probability)
    monkeypatch.setattr(DummyRegressor, "predict", ceiling)
    with pytest.raises(EvidenceError, match="feature order mismatch"):
        evaluate_credit_ml(synthetic_metrics(), str(directory), mode="demo")
    probability.assert_not_called()
    ceiling.assert_not_called()


@pytest.fixture
def reviewed_live_assets(tmp_path, monkeypatch):
    # Hash-pinned byte stand-ins and mocked loaders exercise the live boundary;
    # these are not production pickles, original feature engineering, or labels.
    names = ["monthly_txn_volume", "avg_txn_size", "business_tenure_months",
             "tracking_reliability_pct", "payment_consistency_score", "monthly_revenue",
             "num_employees", "industry_sector_manufacturing", "industry_sector_services"]
    root = tmp_path / "reviewed-repository"
    directory = root / "ml" / "models"
    directory.mkdir(parents=True)
    hashes = {}
    for relative in ("ml/feature_engineering.py", "ml/predict.py",
                     "ml/models/credit_classifier.pkl", "ml/models/credit_regressor.pkl"):
        path = root / relative
        data = b"synthetic reviewed unit-test stand-in; never executed or deserialized"
        path.write_bytes(data)
        key = path.name if path.suffix == ".pkl" else relative
        hashes[key] = hashlib.sha256(data).hexdigest()
    manifest = {"is_demo": False, "sklearn_version": sklearn.__version__,
                "feature_names": names, "default_class": "default", "sha256": hashes}
    manifest_path = directory / "model_manifest.json"
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    frame = pd.DataFrame([[100, 1000, 24, 99, 90, 100000, 5, 1, 0]], columns=names)
    classifier = SimpleNamespace(classes_=np.array(["paid", "default"]),
                                 feature_names_in_=np.array(names))
    regressor = SimpleNamespace(feature_names_in_=np.array(names))
    build = Mock(return_value=frame)
    infer = Mock(return_value={"default_probability": "0.05", "financing_ceiling": "50000"})
    loader = Mock(side_effect=[classifier, regressor])
    modules = {"out_in_reviewed_features": SimpleNamespace(build_features=build),
               "out_in_reviewed_prediction": SimpleNamespace(infer_credit=infer)}
    module_loader = Mock(side_effect=lambda path, name: modules[name])
    monkeypatch.setenv("OUT_IN_REPO_ROOT", str(root))
    monkeypatch.setattr(ml_underwrite_tool.pickle, "load", loader)
    monkeypatch.setattr(ml_underwrite_tool, "_module", module_loader)
    return SimpleNamespace(root=root, directory=directory, manifest=manifest, manifest_path=manifest_path,
                           names=names, frame=frame, classifier=classifier, regressor=regressor,
                           build=build, infer=infer, loader=loader, module_loader=module_loader, modules=modules)


def test_reviewed_live_manifest_schema_is_independent_of_policy_features(reviewed_live_assets):
    case = reviewed_live_assets
    metrics = synthetic_metrics(requested="60000.00")
    metrics["external_features"] = {
        "business_tenure_months": {"value": "24", "provenance": {"source": "synthetic-registration"}},
        "industry_sector": {"value": "manufacturing", "provenance": {"source": "synthetic-registry"}},
    }
    metrics["monthly_txn_volume"] = {"value": "100", "provenance": "synthetic-ledger"}
    metrics["monthly_inflow"] = Decimal("100000.00")
    metrics["monthly_outflow"] = 30000
    original = deepcopy(metrics)
    result = evaluate_credit_ml(metrics, str(case.directory))
    normalized = {**metrics, **{key: str(Decimal(str(metrics[key]))) for key in FEATURES}}
    case.build.assert_called_once_with(normalized)
    assert metrics == original
    supplied = case.build.call_args.args[0]
    assert supplied["provenance"] == original["provenance"]
    assert "monthly_revenue" not in supplied
    assert "tracking_reliability_pct" not in supplied["external_features"]
    assert list(case.frame.columns) == case.names != FEATURES
    assert [call.args[0] for call in case.module_loader.call_args_list] == [
        case.root / "ml/feature_engineering.py", case.root / "ml/predict.py"]
    case.infer.assert_called_once_with(case.classifier, case.regressor, case.frame, "default")
    assert result["passed"] is True
    assert result["is_demo_model"] is False
    assert result["approved_limit"] == "50000.00"
    assert Decimal(result["existing_dsr"]) == Decimal("0.01")
    assert Decimal(result["dsr"]) <= Decimal("0.45")
    assert result["feature_provenance"] == metrics["provenance"]
    assert result["model_hashes"] == {name: case.manifest["sha256"][name]
                                     for name in ("credit_classifier.pkl", "credit_regressor.pkl")}
    assert result["model_manifest_sha256"] == hashlib.sha256(case.manifest_path.read_bytes()).hexdigest()


@pytest.mark.parametrize("defect", ["missing", "extra", "reordered", "duplicate", "empty", "two_rows",
                                  "not_dataframe", "nan", "inf", "negative_inf", "nullable_missing",
                                  "string", "boolean", "complex"])
def test_reviewed_live_adapter_invalid_frames_fail_before_inference(reviewed_live_assets, defect):
    case = reviewed_live_assets
    frame = case.frame.copy()
    if defect == "missing":
        frame = frame.drop(columns=case.names[-1])
    elif defect == "extra":
        frame["unreviewed"] = 1
    elif defect == "reordered":
        frame = frame[list(reversed(case.names))]
    elif defect == "duplicate":
        frame.columns = case.names[:-1] + [case.names[0]]
    elif defect == "empty":
        frame = frame.iloc[:0]
    elif defect == "two_rows":
        frame = pd.concat([frame, frame], ignore_index=True)
    elif defect == "not_dataframe":
        frame = SimpleNamespace(columns=case.names, shape=(1, len(case.names)))
    else:
        value = {"nan": np.nan, "inf": np.inf, "negative_inf": -np.inf,
                 "string": "100", "boolean": True, "complex": 1 + 2j}.get(defect)
        frame[case.names[0]] = pd.Series([pd.NA], dtype="Float64") if defect == "nullable_missing" else value
    case.build.return_value = frame
    with pytest.raises(EvidenceError, match="different schema|finite numeric"):
        evaluate_credit_ml(synthetic_metrics(), str(case.directory))
    case.infer.assert_not_called()


def test_reviewed_live_adapter_does_not_get_fabricated_missing_external_fields(reviewed_live_assets):
    case = reviewed_live_assets

    def require_tenure(metrics):
        assert "business_tenure_months" not in metrics
        assert "external_features" not in metrics
        raise EvidenceError("Missing provenance-bearing business tenure evidence")

    case.build.side_effect = require_tenure
    with pytest.raises(EvidenceError, match="Missing provenance-bearing"):
        evaluate_credit_ml(synthetic_metrics(), str(case.directory))
    case.infer.assert_not_called()


@pytest.mark.parametrize("defect", ["missing", "string", "empty", "duplicate", "blank", "nonstring"])
def test_live_manifest_requires_valid_ordered_feature_names_before_unpickle(reviewed_live_assets, defect):
    case = reviewed_live_assets
    if defect == "missing":
        case.manifest.pop("feature_names")
    else:
        case.manifest["feature_names"] = {"string": "monthly_revenue", "empty": [],
                                          "duplicate": ["revenue", "revenue"], "blank": [" "],
                                          "nonstring": [1]}[defect]
    case.manifest_path.write_text(json.dumps(case.manifest), encoding="utf-8")
    with pytest.raises(EvidenceError, match="Manifest feature schema"):
        evaluate_credit_ml(synthetic_metrics(), str(case.directory))
    case.loader.assert_not_called()
    case.module_loader.assert_not_called()
    case.infer.assert_not_called()


@pytest.mark.parametrize("defect", ["classifier_order", "regressor_order", "unknown_default",
                                  "missing_default", "missing_classes"])
def test_live_model_contract_checked_before_adapters_and_inference(reviewed_live_assets, defect):
    case = reviewed_live_assets
    if defect.endswith("_order"):
        model = case.classifier if defect == "classifier_order" else case.regressor
        model.feature_names_in_ = np.array(list(reversed(case.names)))
    elif defect == "unknown_default":
        case.manifest["default_class"] = "unknown"
    elif defect == "missing_default":
        case.manifest.pop("default_class")
    else:
        del case.classifier.classes_
    case.manifest_path.write_text(json.dumps(case.manifest), encoding="utf-8")
    with pytest.raises(EvidenceError, match="feature order mismatch|Default class label is absent"):
        evaluate_credit_ml(synthetic_metrics(), str(case.directory))
    case.module_loader.assert_not_called()
    case.build.assert_not_called()
    case.infer.assert_not_called()


@pytest.mark.parametrize("model_type", ["classifier", "regressor"])
def test_live_models_without_optional_feature_names_metadata_remain_supported(reviewed_live_assets, model_type):
    case = reviewed_live_assets
    delattr(case.classifier if model_type == "classifier" else case.regressor, "feature_names_in_")
    assert evaluate_credit_ml(synthetic_metrics(), str(case.directory))["passed"] is True
    case.infer.assert_called_once()


@pytest.mark.parametrize("asset", ["ml/feature_engineering.py", "ml/predict.py",
                                 "ml/models/credit_classifier.pkl", "ml/models/credit_regressor.pkl"])
@pytest.mark.parametrize("defect", ["missing", "changed", "unreviewed"])
def test_live_asset_hashes_and_sources_remain_required(reviewed_live_assets, asset, defect):
    case = reviewed_live_assets
    path = case.root / asset
    if defect == "missing":
        path.unlink()
    elif defect == "changed":
        path.write_bytes(path.read_bytes() + b"changed")
    else:
        case.manifest["sha256"].pop(path.name if path.suffix == ".pkl" else asset)
        case.manifest_path.write_text(json.dumps(case.manifest), encoding="utf-8")
    with pytest.raises(EvidenceError, match="missing|hash|unreviewed"):
        evaluate_credit_ml(synthetic_metrics(), str(case.directory))
    if path.suffix == ".pkl":
        case.loader.assert_not_called()
    case.module_loader.assert_not_called()
    case.infer.assert_not_called()


@pytest.mark.parametrize("adapter,entry", [("out_in_reviewed_features", "build_features"),
                                         ("out_in_reviewed_prediction", "infer_credit")])
def test_live_adapter_entry_points_must_be_callable(reviewed_live_assets, adapter, entry):
    case = reviewed_live_assets
    setattr(case.modules[adapter], entry, None)
    with pytest.raises(EvidenceError, match="Repository must implement"):
        evaluate_credit_ml(synthetic_metrics(), str(case.directory))
    case.build.assert_not_called()
    case.infer.assert_not_called()


@pytest.mark.parametrize("defect", ["sklearn_version", "demo_model"])
def test_live_version_and_demo_gates_still_precede_unpickle(reviewed_live_assets, defect):
    case = reviewed_live_assets
    case.manifest["sklearn_version" if defect == "sklearn_version" else "is_demo"] = (
        "0.0.0-test-mismatch" if defect == "sklearn_version" else True)
    case.manifest_path.write_text(json.dumps(case.manifest), encoding="utf-8")
    with pytest.raises(EvidenceError, match="scikit-learn version|Synthetic demo models"):
        evaluate_credit_ml(synthetic_metrics(), str(case.directory))
    case.loader.assert_not_called()
    case.infer.assert_not_called()


@pytest.mark.parametrize("field", EXPECTED_FEATURES)
@pytest.mark.parametrize("defect", ["missing", "nan"])
def test_live_financial_policy_evidence_still_required_before_unpickle(reviewed_live_assets, field, defect):
    case = reviewed_live_assets
    metrics = synthetic_metrics()
    if defect == "missing":
        metrics.pop(field)
    else:
        metrics[field] = "NaN"
    with pytest.raises(EvidenceError):
        evaluate_credit_ml(metrics, str(case.directory))
    case.loader.assert_not_called()
    case.build.assert_not_called()
    case.infer.assert_not_called()


@pytest.mark.parametrize("policy", ["default_risk", "dsr", "invoice_cap", "capacity_cap"])
def test_live_nonfinancial_schema_preserves_financial_safety_gates(reviewed_live_assets, policy):
    case = reviewed_live_assets
    metrics = synthetic_metrics()
    if policy == "default_risk":
        case.infer.return_value["default_probability"] = "0.13"
    elif policy == "dsr":
        metrics = synthetic_metrics(debt="50000.00")
    elif policy == "invoice_cap":
        metrics["invoice_total"] = "10000.00"
    else:
        metrics = synthetic_metrics(debt="44900.00")
    result = evaluate_credit_ml(metrics, str(case.directory))
    if policy in ("default_risk", "dsr"):
        assert result["passed"] is False
        assert result["review_required"] is True
        assert result["approved_limit"] == "0.00"
        assert result["haircut_requested_limit"] == "20000.00"
    elif policy == "invoice_cap":
        assert result["approved_limit"] == "8000.00"
    else:
        assert 0 < Decimal(result["proposed_limit"]) < Decimal("1200")
        assert result["approved_limit"] == "0.00" or Decimal(result["dsr"]) <= Decimal("0.45")


@pytest.mark.parametrize("field,value", [("default_probability", "NaN"), ("default_probability", "Infinity"),
                                         ("default_probability", "-0.1"), ("default_probability", "1.1"),
                                         ("financing_ceiling", "NaN"), ("financing_ceiling", "Infinity"),
                                         ("financing_ceiling", "-1")])
def test_live_invalid_inference_outputs_remain_fail_closed(reviewed_live_assets, field, value):
    case = reviewed_live_assets
    case.infer.return_value[field] = value
    with pytest.raises(EvidenceError, match="nonfinite or invalid"):
        evaluate_credit_ml(synthetic_metrics(), str(case.directory))


@pytest.mark.parametrize("mode", ["", "Demo", "LIVE", "invalid", None, True, 1, [], {}])
def test_invalid_ml_modes_fail_before_unpickle(model_factory, monkeypatch, mode):
    directory, _ = model_factory()
    loader = Mock(side_effect=AssertionError("Invalid modes must fail before unpickle"))
    monkeypatch.setattr(ml_underwrite_tool.pickle, "load", loader)
    with pytest.raises(EvidenceError, match="Mode must be live or demo"):
        evaluate_credit_ml(synthetic_metrics(), str(directory), mode=mode)
    loader.assert_not_called()


@pytest.mark.parametrize("term", [0, 121, -1, 12.5, "12.5", True, None, "NaN", "invalid"])
def test_invalid_or_fractional_loan_duration_fails_before_unpickle(model_factory, monkeypatch, term):
    directory, _ = model_factory()
    loader = Mock(side_effect=AssertionError("Invalid duration must fail before unpickle"))
    monkeypatch.setattr(ml_underwrite_tool.pickle, "load", loader)
    with pytest.raises(EvidenceError):
        evaluate_credit_ml(synthetic_metrics(), str(directory), mode="demo", term_months=term)
    loader.assert_not_called()


@pytest.mark.parametrize("term", [12, "12", Decimal("12")])
def test_integer_loan_duration_preserves_decimal_api_semantics(model_factory, term):
    directory, _ = model_factory()
    result = evaluate_credit_ml(synthetic_metrics(), str(directory), mode="demo", term_months=term)
    assert result["passed"] is True
    assert result["approved_limit"] == "40000.00"


@pytest.fixture
def signed_bundle(tmp_path, monkeypatch, valid_report):
    job = tmp_path / "workspace" / "jobs" / "synthetic-signature-job"
    job.mkdir(parents=True)

    def compile_pdf(loan_terms, template_path, output_path, evidence_hash):
        payload = b"synthetic final PDF byte fixture; rendering intentionally mocked"
        Path(output_path).write_bytes(payload)
        return {"sha256": hashlib.sha256(payload).hexdigest()}

    def ledger(invoices, bank_rows, loan_terms, output_path):
        Path(output_path).write_bytes(b"synthetic ledger byte fixture; rendering intentionally mocked")

    monkeypatch.setattr(artifact_tool, "compile_tnc_agreement", compile_pdf)
    monkeypatch.setattr(artifact_tool, "build_master_ledger", ledger)
    terms = {"principal": "50.00", "loan_cap": "50.00", "annual_rate": "0.08", "term_months": 12}
    source_bytes = b"synthetic source evidence"
    (job / "synthetic-source.txt").write_bytes(source_bytes)
    sources = [{"name": "synthetic-source.txt", "sha256": hashlib.sha256(source_bytes).hexdigest()}]
    names = artifact_tool.build_artifacts(job, sources, [synthetic_invoice()], valid_report,
                                         {"passed": True, "approved_limit": "50.00"}, terms, "demo")
    return job, job / "artifacts", names


def public_key(bundle, filename):
    return serialization.load_pem_public_key((bundle / filename).read_bytes())


def test_manifest_signature_hashes_and_role_attestations_verify_independently(signed_bundle):
    job, bundle, names = signed_bundle
    manifest_bytes = (bundle / "verification_manifest.json").read_bytes()
    manifest = json.loads(manifest_bytes)
    canonical = json.dumps(manifest, sort_keys=True, separators=(",", ":"), ensure_ascii=False,
                           allow_nan=False).encode("utf-8")
    assert manifest_bytes == canonical
    public_key(bundle, "demo_public_key.pem").verify((bundle / "verification_manifest.sig").read_bytes(), canonical)
    for entry in manifest["outputs"]:
        assert hashlib.sha256((bundle / entry["name"]).read_bytes()).hexdigest() == entry["sha256"]
    assert set(names) == {path.name for path in bundle.iterdir()}
    assert manifest["real_funds"] is False
    assert manifest["legal_status"] == "demo_attestation_not_legal_countersignature"
    assert not manifest["policy_gates"]["computed"]["legal_countersignature_claim"]
    assert manifest["sources"] == [{"name": "synthetic-source.txt", "sha256": hashlib.sha256(
        (job / "synthetic-source.txt").read_bytes()).hexdigest()}]
    pdf = (bundle / "Out_In_TNC_Agreement.pdf").read_bytes()
    keys = []
    for role in ("borrower", "lender"):
        key = public_key(bundle, f"demo_{role}_public_key.pem")
        keys.append(key.public_bytes(serialization.Encoding.Raw, serialization.PublicFormat.Raw))
        message = b"OUT_IN_DEMO_PDF_ATTESTATION_V1\x00" + role.encode() + b"\x00" + pdf
        key.verify((bundle / f"{role}_attestation.sig").read_bytes(), message)
    keys.append(public_key(bundle, "demo_public_key.pem").public_bytes(
        serialization.Encoding.Raw, serialization.PublicFormat.Raw))
    assert len(set(keys)) == 3
    assert not any("private" in path.name.lower() or b"PRIVATE KEY" in path.read_bytes()
                   for path in bundle.iterdir())
    assert (job.parent.parent / "keys").is_dir()


@pytest.mark.parametrize("tamper", ["manifest", "pdf", "role", "key", "signature"])
def test_detached_signatures_reject_tampered_bytes_role_and_key(signed_bundle, tamper):
    _, bundle, _ = signed_bundle
    if tamper == "manifest":
        signature = (bundle / "verification_manifest.sig").read_bytes()
        payload = json.loads((bundle / "verification_manifest.json").read_bytes())
        payload["risk"]["approved_limit"] = "50.01"
        message = json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()
        key = public_key(bundle, "demo_public_key.pem")
    else:
        signature = (bundle / "borrower_attestation.sig").read_bytes()
        pdf = (bundle / "Out_In_TNC_Agreement.pdf").read_bytes()
        role = b"lender" if tamper == "role" else b"borrower"
        message = b"OUT_IN_DEMO_PDF_ATTESTATION_V1\x00" + role + b"\x00" + pdf
        key = public_key(bundle, "demo_lender_public_key.pem" if tamper == "key" else "demo_borrower_public_key.pem")
        if tamper == "pdf":
            message += b"synthetic tampered byte"
        elif tamper == "signature":
            signature = bytes([signature[0] ^ 1]) + signature[1:]
    with pytest.raises(InvalidSignature):
        key.verify(signature, message)


@pytest.mark.parametrize("name", ["../private.pem", "/absolute.pdf", "C:/private.pem", "nested/../../source.pdf",
                                 "nested\\source.pdf"])
def test_artifact_source_path_traversal_rejected_before_any_output(tmp_path, valid_report, name):
    job = tmp_path / "workspace" / "jobs" / "synthetic-job"
    job.mkdir(parents=True)
    with pytest.raises(ValueError, match="relative|traversal"):
        artifact_tool.build_artifacts(job, [{"name": name, "sha256": "0" * 64}], [synthetic_invoice()],
                                      valid_report, {"passed": True},
                                      {"principal": "50.00", "annual_rate": "0.08", "term_months": 12}, "demo")
    assert not (job / "staging").exists()
    assert not (job / "artifacts").exists()


@pytest.mark.parametrize("identifier", ["=1+1", "+1+1", "-1+1", "@SUM(1,1)"])
def test_exported_untrusted_invoice_ids_remain_literal_not_spreadsheet_formulas(tmp_path, identifier):
    from openpyxl import load_workbook
    invoice = synthetic_invoice(invoice_id=identifier)
    row = synthetic_bank_row(invoice_id=identifier)
    path = tmp_path / "synthetic-ledger.xlsx"
    artifact_tool.build_master_ledger([invoice], [row],
                                     {"principal": "50.00", "annual_rate": "0.08", "term_months": 12}, str(path))
    book = load_workbook(path, data_only=False)
    try:
        for sheet, coordinate in (("InvoiceLines", "A2"), ("Invoices", "A2"), ("Bank", "D2")):
            assert book[sheet][coordinate].value == identifier
            assert book[sheet][coordinate].data_type == "s"
        audit_text = "".join(row[0] for row in book["DecimalAudit"].iter_rows(min_row=2, min_col=2, max_col=2,
                                                                          values_only=True) if row[0])
        assert json.loads(audit_text)["invoices"][0]["invoice_id"] == identifier
    finally:
        book.close()


def synthetic_risk(passed=True):
    base = {"passed": passed, "review_required": not passed, "default_probability": "0.05",
            "dsr": "0.35", "existing_dsr": "0.01", "financing_ceiling": "65000.00",
            "requested_limit": "40000.00", "approved_limit": "40000.00" if passed else "0.00",
            "proposed_limit": "40000.00" if passed else "20000.00",
            "haircut_requested_limit": None if passed else "20000.00",
            "gates": {"probability_le_0_12": True, "all_in_dsr_le_0_45": passed},
            "is_demo_model": True}
    if not passed:
        base["gates"]["all_in_dsr_le_0_45"] = False
    return base


def test_decision_memo_copies_verbatim_figures_and_explains_approval(valid_report):
    risk = synthetic_risk(passed=True)
    terms = {"requested_limit": "40000.00"}
    memo = build_decision_memo(valid_report, risk, terms)
    assert memo["is_demo"] is True and memo["outcome"] == "APPROVED"
    assert memo["narrative_source"] == "deterministic"
    figures = expected_figures(valid_report, risk, terms)
    assert memo["figures"] == figures
    assert figures["variance"] == valid_report["variance"]
    assert figures["invoice_total"] == valid_report["invoice_total"]
    assert figures["verified_inflow"] == valid_report["verified_inflow"]
    assert figures["approved_limit"] == "40000.00" and figures["default_probability"] == "0.05"
    assert memo["reasons"] and any("0.12" in reason for reason in memo["reasons"])
    assert any("0.45" in reason for reason in memo["reasons"])
    assert "no financing authority" in memo["authority"]
    assert json.dumps(memo, allow_nan=False)


def test_decision_memo_review_path_records_bounded_proposal_not_approval(valid_report):
    risk = synthetic_risk(passed=False)
    memo = build_decision_memo(valid_report, risk, {"requested_limit": "40000.00"})
    assert memo["outcome"] == "REVIEW_REQUIRED"
    assert memo["figures"]["approved_limit"] == "0.00"
    assert memo["figures"]["proposed_limit"] == "20000.00"
    assert any("never an approval" in reason for reason in memo["reasons"])
    assert any("review" in step.lower() for step in memo["next_steps"])


def test_decision_memo_tolerates_partial_risk_evidence_without_inventing_figures(valid_report):
    risk = {"passed": True, "approved_limit": "40000.00", "proposed_limit": "40000.00"}
    memo = build_decision_memo(valid_report, risk, {"requested_limit": "40000.00"})
    assert "dsr" not in memo["figures"] and "default_probability" not in memo["figures"]
    assert memo["reasons"]


def test_decision_memo_rejects_narrative_with_altered_figures(valid_report):
    risk = synthetic_risk(passed=True)
    terms = {"requested_limit": "40000.00"}
    figures = expected_figures(valid_report, risk, terms)
    narrative = {"headline": "Synthetic altered narrative", "summary": "Tampered figure attempt.",
                 "reasons": ["Synthetic reason"], "next_steps": [],
                 "figures": {**figures, "approved_limit": "999999.00"}}
    with pytest.raises(EvidenceError, match="figures differ"):
        build_decision_memo(valid_report, risk, terms, narrative=narrative, narrative_source="sdk_agent")


def test_decision_memo_accepts_narrative_with_exact_figures(valid_report):
    risk = synthetic_risk(passed=True)
    terms = {"requested_limit": "40000.00"}
    figures = expected_figures(valid_report, risk, terms)
    narrative = {"headline": "Synthetic plain-language headline", "summary": "Synthetic summary.",
                 "reasons": ["Synthetic reason one."], "next_steps": ["Synthetic step."],
                 "figures": figures}
    memo = build_decision_memo(valid_report, risk, terms, narrative=narrative, narrative_source="sdk_agent")
    assert memo["headline"] == "Synthetic plain-language headline"
    assert memo["figures"] == figures and memo["narrative_source"] == "sdk_agent"


def test_decision_memo_rejects_invalid_evidence_figure(valid_report):
    risk = synthetic_risk(passed=True)
    risk["default_probability"] = 0.05  # Binary float, not an evidence string.
    with pytest.raises(EvidenceError):
        build_decision_memo(valid_report, risk, {"requested_limit": "40000.00"})


def test_decision_memo_accepts_long_decimal_and_scientific_evidence_strings(valid_report):
    risk = synthetic_risk(passed=True)
    risk["default_probability"] = "0.041666666666666664"
    risk["dsr"] = "0.3433333333333333333333333333"
    memo = build_decision_memo(valid_report, risk, {"requested_limit": "40000.00"})
    assert memo["figures"]["default_probability"] == "0.041666666666666664"
    assert memo["figures"]["dsr"] == "0.3433333333333333333333333333"


def test_bundle_selfcheck_verifies_real_committed_bundle(signed_bundle):
    job, bundle, names = signed_bundle
    result = verify_bundle_integrity(job)
    assert result["verified"] is True and result["self_check"] == "post_commit"
    assert result["outputs_checked"] == 7 and result["sources_checked"] == 1
    assert result["real_funds"] is False and result["is_demo"] is True


def test_bundle_selfcheck_fails_closed_on_tampered_bundle(signed_bundle):
    job, bundle, names = signed_bundle
    pdf = bundle / "Out_In_TNC_Agreement.pdf"
    original = pdf.read_bytes()
    try:
        pdf.write_bytes(bytes([original[0] ^ 1]) + original[1:])
        with pytest.raises(EvidenceError, match="self-verification"):
            verify_bundle_integrity(job)
    finally:
        pdf.write_bytes(original)


def test_bundle_selfcheck_fails_closed_when_bundle_missing(tmp_path):
    with pytest.raises(EvidenceError, match="missing"):
        verify_bundle_integrity(tmp_path / "synthetic-job")


def test_pdfium_repairs_hold_process_lock_but_initial_text_extraction_does_not(monkeypatch):
    import threading
    from tools import ocr_tool

    lock = threading.RLock()
    observations = []
    sentinel = object()

    def normalize(path, regions, cycle):
        acquired = []
        def probe():
            available = lock.acquire(blocking=False)
            acquired.append(available)
            if available:
                lock.release()
        worker = threading.Thread(target=probe)
        worker.start()
        worker.join(timeout=5)
        assert not worker.is_alive()
        observations.append((path, regions, cycle, acquired))
        return sentinel

    monkeypatch.setattr(ocr_tool, "_PDFIUM_LOCK", lock)
    monkeypatch.setattr(ocr_tool, "_normalize_pdf", normalize)
    assert ocr_tool.run_ocr_normalizer("synthetic.pdf") is sentinel
    for cycle in (1, 2, 3):
        assert ocr_tool.run_ocr_normalizer("synthetic.pdf", [], cycle) is sentinel
    assert [row[3] for row in observations] == [[True], [False], [False], [False]]
    assert [row[2] for row in observations] == [0, 1, 2, 3]
